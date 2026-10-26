"""`scripts/clean_cert_validity.py` 的验收测试。

这个脚本会**改写 Excel**，所以测试必须守住三件事：
1. 区间取**右端**（到期日），不是左端；
2. 端点不合法时**原值不动**（绝不拿左端顶替 —— 见 AGENTS.md 约定 14）；
3. **只动目标列**，其余列/行数/表头一律不变，且重复执行无副作用（幂等）。
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(ROOT / "scripts"))

from clean_cert_validity import (  # noqa: E402
    DEFAULT_COLUMN,
    MODE_ALL,
    MODE_RANGE,
    ST_CLEANED,
    ST_CLEANED_EST,
    ST_KEEP_MONTH,
    ST_KEEP_SINGLE,
    ST_UNPARSABLE,
    assert_output_outside_inbound,
    clean_workbook,
    decide,
    default_paths,
)

SHEET = "小黄卡清单"
HEADERS = ["姓名", "身份证号码", "公司名称", "证书编号", "证书开始日期", "证书有效期", "日期来源"]


# --------------------------------------------------------------- 纯函数：decide
def test_range_takes_right_end():
    """用户明确要求：2026/09/18-2028/09/30 -> 2028/09/30。"""
    d = decide("2026/09/18-2028/09/30", MODE_RANGE, "%Y/%m/%d")
    assert d.new_value == "2028/09/30"
    assert d.status == ST_CLEANED
    assert d.shape == "区间"


def test_range_with_illegal_right_end_stays_untouched():
    """右端 4 月 31 日不存在 —— 不猜、不拿左端顶替，原值保持不动。"""
    d = decide("2026/04/17-2028/04/31", MODE_RANGE, "%Y/%m/%d")
    assert d.new_value is None
    assert d.status == ST_UNPARSABLE
    assert "到期日期不合法" in d.detail


def test_reversed_range_is_not_fixed():
    d = decide("2028/09/30-2026/09/18", MODE_RANGE, "%Y/%m/%d")
    assert d.new_value is None
    assert d.status == ST_UNPARSABLE


def test_single_date_kept_by_default():
    d = decide("2027/03/31", MODE_RANGE, "%Y/%m/%d")
    assert d.new_value is None
    assert d.status == ST_KEEP_SINGLE


def test_year_month_kept_by_default_and_estimated_in_all_mode():
    """只有年月时，按月末取值是**估计值**，必须显式开启。"""
    kept = decide("2027年1月", MODE_RANGE, "%Y/%m/%d")
    assert kept.new_value is None
    assert kept.status == ST_KEEP_MONTH

    folded = decide("2027年1月", MODE_ALL, "%Y/%m/%d")
    assert folded.new_value == "2027/01/31"
    assert folded.status == ST_CLEANED_EST
    assert "估计值" in folded.detail


def test_empty_and_garbage_are_unparsable():
    for raw in (None, "", "   ", "待定", "2026/13/45-2028/01/01"):
        d = decide(raw, MODE_ALL, "%Y/%m/%d")
        assert d.new_value is None
        assert d.status == ST_UNPARSABLE


# --------------------------------------------------------------- 端到端：改文件
def _make_workbook(path: Path, validity_values: list[object]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(HEADERS)
    for i, value in enumerate(validity_values, start=1):
        ws.append([f"人{i}", f"3101011990010{i:04d}", "某公司", f"C{i}", "2026/09/18", value, "证面原文"])
    wb.save(path)
    wb.close()
    return path


def _read_column(path: Path, column: str = DEFAULT_COLUMN) -> tuple[list, list[list]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[SHEET]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    idx = rows[0].index(column)
    return [r[idx] for r in rows[1:]], rows


def test_clean_workbook_only_touches_target_column(tmp_path: Path):
    values = [
        "2026/09/18-2028/09/30",   # 区间 → 折叠
        "2027/03/31",              # 单日期 → 不动
        "2027年1月",                # 年月 → 不动（range 模式）
        "2026/04/17-2028/04/31",   # 右端非法 → 不动
        2027,                      # 非日期垃圾 → 不动
    ]
    src = _make_workbook(tmp_path / "in.xlsx", values)

    wb, ws, summary, rows = clean_workbook(src, SHEET, DEFAULT_COLUMN, MODE_RANGE, "%Y/%m/%d")
    assert wb is not None
    out = tmp_path / "out.xlsx"
    wb.save(out)
    wb.close()

    assert summary.total == len(values)
    assert summary.changed == 1
    assert summary.by_status[ST_CLEANED] == 1
    assert summary.by_status[ST_UNPARSABLE] == 2  # 右端非法 + 非日期垃圾

    after, rows_after = _read_column(out)
    # 非日期的离谱值（int 2027）也必须原样保留 —— 不解释、不纠错
    assert after == ["2028/09/30", "2027/03/31", "2027年1月", "2026/04/17-2028/04/31", 2027]

    # 只动目标列：其余列与行数必须完全一致
    _, rows_before = _read_column(src)
    assert len(rows_before) == len(rows_after)
    assert rows_before[0] == rows_after[0]
    for before_row, after_row in zip(rows_before[1:], rows_after[1:]):
        idx = HEADERS.index(DEFAULT_COLUMN)
        assert [v for i, v in enumerate(before_row) if i != idx] == [
            v for i, v in enumerate(after_row) if i != idx
        ]


def test_second_run_is_idempotent(tmp_path: Path):
    src = _make_workbook(tmp_path / "in.xlsx", ["2026/09/18-2028/09/30", "2026/04/17-2028/04/31"])

    wb, _, first, _ = clean_workbook(src, SHEET, DEFAULT_COLUMN, MODE_RANGE, "%Y/%m/%d")
    once = tmp_path / "once.xlsx"
    wb.save(once)
    wb.close()
    assert first.changed == 1

    wb2, _, second, _ = clean_workbook(once, SHEET, DEFAULT_COLUMN, MODE_RANGE, "%Y/%m/%d")
    wb2.close()
    assert second.changed == 0, "重复执行不应再改任何值"


def test_dry_run_does_not_load_write_target(tmp_path: Path):
    src = _make_workbook(tmp_path / "in.xlsx", ["2026/09/18-2028/09/30"])
    wb, _, summary, rows = clean_workbook(
        src, SHEET, DEFAULT_COLUMN, MODE_RANGE, "%Y/%m/%d", apply=False
    )
    assert wb is None, "dry-run 不应返回可写 workbook"
    assert summary.changed == 1
    assert rows and rows[0].after == "2028/09/30"
    # 源文件必须没被改动
    assert _read_column(src)[0] == ["2026/09/18-2028/09/30"]


def test_missing_column_reports_available_headers(tmp_path: Path):
    src = _make_workbook(tmp_path / "in.xlsx", ["2027/03/31"])
    with pytest.raises(SystemExit) as exc:
        clean_workbook(src, SHEET, "不存在的列", MODE_RANGE, "%Y/%m/%d")
    assert "找不到列" in str(exc.value)


# --------------------------------------------------------------- 输出位置护栏
def test_output_defaults_to_cleaned_dir_not_inbound():
    """清洗结果不能落在 data/inbound/ 下（否则入站会把两份台账互相覆盖）。"""
    inbound = ROOT / "data" / "inbound" / "certificates" / "小黄卡简表.xlsx"
    output, report = default_paths(inbound)
    assert output == ROOT / "data" / "cleaned" / "certificates" / "小黄卡简表.cleaned.xlsx"
    assert report.name == "小黄卡简表.cleaned.report.csv"
    # 落在入站目录外，护栏不拦
    assert_output_outside_inbound(output)


def test_output_inside_inbound_is_rejected():
    bad = ROOT / "data" / "inbound" / "certificates" / "x.cleaned.xlsx"
    with pytest.raises(SystemExit) as exc:
        assert_output_outside_inbound(bad)
    assert "data/inbound" in str(exc.value)


def test_output_outside_repo_is_allowed(tmp_path: Path):
    assert_output_outside_inbound(tmp_path / "x.cleaned.xlsx")  # 不抛异常
