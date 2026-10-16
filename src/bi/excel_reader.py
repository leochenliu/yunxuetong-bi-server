"""Excel 读取层：把 .xlsx 读成内存数据集，只负责「读进来 + 保真」。

保真的关键点：
  - 保留每一行的 **Excel 物理行号**（`__row__`）——证据追溯要能指回「原始文件第几行」。
  - 重复列名自动去重（`备注1`, `备注1_2`），不静默丢列。
  - 全空行跳过（Excel 尾部常有）。
  - 日期保持 datetime 对象，不在这一层做业务解释（那是 derive 的职责）。

实现说明：使用 openpyxl 只读模式 + 单次流式遍历。只读模式下**不支持随机单元格
访问**（`ws.cell()`），因此表头也通过同一次遍历取得。
"""

from __future__ import annotations

import datetime as _dt
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterator, Sequence

import openpyxl

ROW_NUMBER_KEY = "__row__"


@dataclass
class Dataset:
    name: str
    sheet: str
    columns: list[str]
    rows: list[dict[str, Any]]
    source_file: str
    notes: list[str] = field(default_factory=list)

    def __len__(self) -> int:
        return len(self.rows)

    def column_values(self, column: str) -> list[Any]:
        return [r.get(column) for r in self.rows]


def _normalize_cell(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, str):
        t = v.strip()
        return t if t != "" else None
    return v


def _dedup_columns(header_values: Sequence[Any]) -> list[tuple[str, int]]:
    """返回 [(列名, 列下标1-based)]，重复列名自动加后缀。"""
    seen: dict[str, int] = {}
    cols: list[tuple[str, int]] = []
    for i, raw in enumerate(header_values, start=1):
        v = _normalize_cell(raw)
        name = "" if v is None else str(v).strip()
        if name == "":
            name = f"column_{i}"
        n = seen.get(name, 0)
        seen[name] = n + 1
        cols.append((name if n == 0 else f"{name}_{n + 1}", i))
    return cols


def iter_rows(
    path: str | Path,
    sheet: str | None = None,
    header_row: int = 1,
) -> Iterator[Dataset]:
    """按工作表逐个产出数据集（流式，适合大表）。"""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Excel 文件不存在: {path}")

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            if sheet and ws.title != sheet:
                continue
            cols: list[tuple[str, int]] | None = None
            rows: list[dict[str, Any]] = []
            for r_idx, raw_row in enumerate(ws.iter_rows(values_only=True), start=1):
                if r_idx < header_row:
                    continue
                if r_idx == header_row:
                    cols = _dedup_columns(raw_row)
                    continue
                if not cols:
                    break
                obj: dict[str, Any] = {ROW_NUMBER_KEY: r_idx}
                any_value = False
                n_cells = len(raw_row)
                for name, idx in cols:
                    v = _normalize_cell(raw_row[idx - 1]) if idx - 1 < n_cells else None
                    obj[name] = v
                    if v is not None:
                        any_value = True
                if not any_value:
                    continue  # 全空行
                rows.append(obj)
            if not cols:
                continue
            yield Dataset(
                name=path.stem,
                sheet=ws.title,
                columns=[c[0] for c in cols],
                rows=rows,
                source_file=str(path),
            )
    finally:
        wb.close()


def read_workbook(
    path: str | Path,
    sheet: str | None = None,
    header_row: int = 1,
) -> list[Dataset]:
    return list(iter_rows(path, sheet=sheet, header_row=header_row))


def read_primary(path: str | Path, sheet: str | None = None) -> Dataset:
    """读第一个（或指定）工作表——本项目的 Excel 只有一个「总表」。"""
    datasets = read_workbook(path, sheet=sheet)
    if not datasets:
        raise ValueError(f"未读到任何工作表{'(sheet=' + sheet + ')' if sheet else ''}: {path}")
    return datasets[0]


def file_fingerprint(path: str | Path) -> dict[str, Any]:
    """文件指纹，用于上传版本去重与溯源。"""
    import hashlib

    p = Path(path)
    h = hashlib.sha256()
    with p.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    st = p.stat()
    return {
        "sha256": h.hexdigest(),
        "size_bytes": st.st_size,
        "modified_at": _dt.datetime.fromtimestamp(st.st_mtime).isoformat(timespec="seconds"),
        "filename": p.name,
    }
